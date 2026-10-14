"""Streaming file staging. Business records are never changed here."""
import csv
import hashlib
import io
from pathlib import Path
from openpyxl import load_workbook
from .models import ImportRow


class UploadError(ValueError):
    pass


def checksum(upload):
    digest = hashlib.sha256()
    for chunk in upload.chunks():
        digest.update(chunk)
    upload.seek(0)
    return digest.hexdigest()


def _csv_records(path):
    try:
        binary = open(path, "rb")
        text = io.TextIOWrapper(binary, encoding="utf-8-sig", newline="")
        sample = text.read(4096)
        text.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(text, dialect=dialect, strict=True)
        headers = next(reader, None)
        if not headers:
            raise UploadError("The file has no header row.")
        yield headers, None, None
        while True:
            source_row = reader.line_num + 1
            try:
                values = next(reader)
            except StopIteration:
                break
            if not values or all(not str(value).strip() for value in values):
                continue
            if len(values) != len(headers):
                raise UploadError(f"Source row {source_row} has {len(values)} values; expected {len(headers)}.")
            yield None, source_row, [str(v) for v in values]
    except UnicodeDecodeError as exc:
        raise UploadError("CSV must be UTF-8 encoded. Save it as UTF-8 and upload again.") from exc
    except csv.Error as exc:
        raise UploadError(f"Malformed CSV: {exc}") from exc
    finally:
        try:
            text.close()
        except UnboundLocalError:
            pass


def _xlsx_records(path):
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        headers = [str(v).strip() if v is not None else "" for v in next(iterator, ())]
        if not headers:
            raise UploadError("The workbook has no header row.")
        yield headers, None, None
        for number, values in enumerate(iterator, start=2):
            if not any(v is not None and str(v).strip() for v in values):
                continue
            if len(values) > len(headers) and any(v is not None for v in values[len(headers):]):
                raise UploadError(f"Source row {number} contains values beyond the header columns.")
            yield None, number, ["" if v is None else str(v) for v in values[:len(headers)]] + [""] * max(0, len(headers) - len(values))
    except Exception as exc:
        if isinstance(exc, UploadError):
            raise
        raise UploadError(f"Could not read XLSX workbook: {exc}") from exc
    finally:
        if "workbook" in locals():
            workbook.close()


def stage_file(job):
    suffix = Path(job.filename).suffix.lower()
    if suffix not in {".csv", ".xlsx"}:
        raise UploadError("Choose a CSV or XLSX file.")
    records = _csv_records(job.source_file.path) if suffix == ".csv" else _xlsx_records(job.source_file.path)
    headers, _, _ = next(records)
    headers = [str(h).strip() for h in headers]
    if any(not h for h in headers):
        raise UploadError("Every source column needs a header.")
    if len(set(h.casefold() for h in headers)) != len(headers):
        raise UploadError("The file contains duplicate column headers.")
    batch = []
    count = 0
    for _, source_row, values in records:
        batch.append(ImportRow(job=job, source_row=source_row, original=dict(zip(headers, values))))
        count += 1
        if len(batch) >= 500:
            ImportRow.objects.bulk_create(batch)
            batch.clear()
    if batch:
        ImportRow.objects.bulk_create(batch)
    if not count:
        raise UploadError("The file has no data records.")
    job.source_headers = headers
    job.total_rows = count
    job.save(update_fields=["source_headers", "total_rows"])
    return headers, count
